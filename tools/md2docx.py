"""Markdown -> .docx for the issue record files in this repository.

Every disposition, decision record and review in this repo is written as
markdown first. The board reads .docx, so each one needs a paired .docx.
This renders the subset of markdown those files actually use:

  # / ## / ### / ####   headings
  paragraphs            with **bold**, *italic*, `code`, [text](link)
  - / * bullets         one level
  1. ordered lists      one level
  | a | b |             pipe tables with a --- separator row
  ---                   horizontal rule
  > blockquote          rendered as an indented callout

House style matches AgentFilesV2/render.py: Calibri 10.5pt body, blue
headings, "Light Grid Accent 1" tables.

Usage:
    python3 tools/md2docx.py INPUT.md [OUTPUT.docx]
    python3 tools/md2docx.py --all          # every *.md lacking a .docx
"""
import re
import sys
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROW = re.compile(r"^\s*\|(.+)\|\s*$")
SEP = re.compile(r"^\s*\|[\s:|-]+\|\s*$")
MARKER = re.compile(r"^\s*([-*]|\d+[.)])\s+")

# Inline spans. These nest — a bold run in these documents routinely contains
# an italic quotation, a code span or a board link — so they are matched one at
# a time and the contents re-scanned rather than consumed in a single pass.
SPANS = (
    ("bold", re.compile(r"\*\*(.+?)\*\*", re.S)),
    ("code", re.compile(r"`([^`]+)`")),
    ("link", re.compile(r"\[([^\]]*)\]\(([^)]*)\)")),
    # An italic quotation often contains a bold phrase, so its contents may
    # span ** pairs but never a single stray *. The delimiters may not sit
    # against whitespace, so a lone "*" in prose is left alone.
    ("italic", re.compile(r"(?<!\*)\*(?!\s)((?:[^*\n]|\*\*)+?)(?<!\s)\*(?!\*)")),
)


def _is_block_start(line):
    """True if this line opens a block, so a paragraph or list item ends here."""
    s = line.strip()
    return (not s or s.startswith("#") or s.startswith(">")
            or s in ("---", "***", "___")
            or bool(ROW.match(line.rstrip())) or bool(MARKER.match(line)))


def _cells(line):
    """Split one pipe-table row, honouring \\| escapes."""
    body = ROW.match(line).group(1)
    out, cur, i = [], "", 0
    while i < len(body):
        if body[i] == "\\" and i + 1 < len(body) and body[i + 1] == "|":
            cur += "|"
            i += 2
        elif body[i] == "|":
            out.append(cur.strip())
            cur = ""
            i += 1
        else:
            cur += body[i]
            i += 1
    out.append(cur.strip())
    return out


def _plain(paragraph, text, style):
    if not text:
        return
    r = paragraph.add_run(text)
    r.bold = style["bold"]
    r.italic = style["italic"]
    if style["code"]:
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)
    if style["link"]:
        r.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)


def _runs(paragraph, text, style=None):
    """Add inline-formatted runs to a paragraph, honouring nested spans."""
    style = style or {"bold": False, "italic": False, "code": False, "link": False}

    # Take the span that starts earliest, so nesting resolves outside-in.
    best = None
    for kind, pattern in SPANS:
        m = pattern.search(text)
        if m and (best is None or m.start() < best[1].start()):
            best = (kind, m)

    if best is None:
        _plain(paragraph, text, style)
        return

    kind, m = best
    _plain(paragraph, text[:m.start()], style)
    # A link's target is an internal board URL that means nothing on paper, so
    # only its words are kept.
    _runs(paragraph, m.group(1), {**style, kind: True})
    _runs(paragraph, text[m.end():], style)


def _blocks(lines):
    """Group markdown lines into (kind, payload) blocks."""
    out, i = [], 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
        elif stripped in ("---", "***", "___"):
            out.append(("hr", None))
            i += 1
        elif stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            out.append((f"h{min(level, 4)}", stripped.lstrip("# ").strip()))
            i += 1
        elif ROW.match(line):
            rows = []
            while i < len(lines) and ROW.match(lines[i].rstrip()):
                if not SEP.match(lines[i].rstrip()):
                    rows.append(_cells(lines[i].rstrip()))
                i += 1
            if rows:
                width = max(len(r) for r in rows)
                out.append(("table", [r + [""] * (width - len(r)) for r in rows]))
        elif MARKER.match(line):
            # These files are hard-wrapped, so an item continues onto the
            # following lines until a blank line or the next block.
            ordered = not re.match(r"^\s*[-*]\s+", line)
            items = []
            while i < len(lines) and MARKER.match(lines[i]):
                item = [MARKER.sub("", lines[i].rstrip())]
                i += 1
                while i < len(lines) and not _is_block_start(lines[i]):
                    item.append(lines[i].strip())
                    i += 1
                items.append(" ".join(item))
            out.append(("ol" if ordered else "ul", items))
        elif stripped.startswith(">"):
            para = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                para.append(lines[i].strip().lstrip("> ").strip())
                i += 1
            out.append(("quote", " ".join(para)))
        else:
            # A paragraph runs until a blank line or the start of another block.
            para = [stripped]
            i += 1
            while i < len(lines) and not _is_block_start(lines[i]):
                para.append(lines[i].strip())
                i += 1
            out.append(("p", " ".join(para)))
    return out


def render(blocks, path):
    d = docx.Document()
    s = d.sections[0]
    s.left_margin = s.right_margin = Inches(0.9)
    s.top_margin = s.bottom_margin = Inches(0.85)
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10.5)

    first_heading = True
    for kind, payload in blocks:
        if kind == "h1" and first_heading:
            p = d.add_paragraph(style="Title")
            _runs(p, payload)
            first_heading = False
        elif kind in ("h1", "h2", "h3", "h4"):
            level = {"h1": 1, "h2": 1, "h3": 2, "h4": 3}[kind]
            p = d.add_paragraph(style=f"Heading {level}")
            _runs(p, payload)
            first_heading = False
        elif kind == "p":
            _runs(d.add_paragraph(), payload)
        elif kind == "quote":
            p = d.add_paragraph()
            _runs(p, payload)
            p.paragraph_format.left_indent = Inches(0.3)
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        elif kind == "ul":
            for item in payload:
                _runs(d.add_paragraph(style="List Bullet"), item)
        elif kind == "ol":
            for item in payload:
                _runs(d.add_paragraph(style="List Number"), item)
        elif kind == "table":
            t = d.add_table(rows=len(payload), cols=len(payload[0]))
            t.style = "Light Grid Accent 1"
            t.alignment = WD_TABLE_ALIGNMENT.LEFT
            for ri, row in enumerate(payload):
                for ci, cell in enumerate(row):
                    para = t.cell(ri, ci).paragraphs[0]
                    _runs(para, cell)
                    for r in para.runs:
                        r.font.size = Pt(9.5)
                        if ri == 0:
                            r.bold = True
            d.add_paragraph()
        elif kind == "hr":
            p = d.add_paragraph("_" * 86)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    d.save(path)
    return path


def convert(src, dest=None):
    src = Path(src)
    dest = Path(dest) if dest else src.with_suffix(".docx")
    render(_blocks(src.read_text(encoding="utf-8").splitlines()), dest)
    return dest


def main(argv):
    if not argv:
        print(__doc__)
        return 1
    if argv[0] == "--all":
        root = Path(argv[1]) if len(argv) > 1 else Path(".")
        todo = [p for p in sorted(root.glob("*.md"))
                if p.name != "README.md" and not p.with_suffix(".docx").exists()]
        if not todo:
            print("nothing to convert")
            return 0
        for p in todo:
            print(f"{p.name} -> {convert(p).name}")
        return 0
    out = convert(argv[0], argv[1] if len(argv) > 1 else None)
    print(f"{argv[0]} -> {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
